# -*- coding: utf-8 -*-
"""
Script gerador do E-Book / Manual do Usuário Completo em formato .DOCX
Sistema: Painel de Gestão
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCX_OUTPUT = r"d:\Projetos\painel_gestao\Ebook_Manual_Painel_de_Gestao.docx"

def set_cell_background(cell, hex_color):
    """Define a cor de fundo de uma célula da tabela"""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Define margens internas de uma célula em twips"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text, title="DICA PRÁTICA", callout_type="tip"):
    """Cria uma caixa de destaque visual com borda colorida e fundo suave"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    
    bg_color = "EFF6FF" if callout_type == "tip" else ("ECFDF5" if callout_type == "success" else "FEF3C7")
    border_color = "2563EB" if callout_type == "tip" else ("059669" if callout_type == "success" else "D97706")
    icon = "💡" if callout_type == "tip" else ("✅" if callout_type == "success" else "⚠️")
    
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    # Borda esquerda grossa
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    '''
    tcPr.append(parse_xml(borders_xml))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    run_title = p.add_run(f"{icon} {title}: ")
    run_title.font.bold = True
    run_title.font.size = Pt(10)
    run_title.font.name = "Arial"
    if callout_type == "tip":
        run_title.font.color.rgb = RGBColor(30, 64, 175)
    elif callout_type == "success":
        run_title.font.color.rgb = RGBColor(4, 120, 87)
    else:
        run_title.font.color.rgb = RGBColor(180, 83, 9)
        
    run_text = p.add_run(text)
    run_text.font.size = Pt(9.5)
    run_text.font.name = "Arial"
    run_text.font.color.rgb = RGBColor(51, 65, 85)
    
    # Espaço após a caixa
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def format_heading(p, text, level):
    """Formata títulos com hierarquia visual moderna"""
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.bold = True
    
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(15, 23, 42) # Slate 900
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(30, 64, 175) # Blue 800
    elif level == 3:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(71, 85, 105) # Slate 600

def add_body_p(doc, text):
    """Adiciona parágrafo de texto padrão"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_bullet_item(doc, bold_prefix, text):
    """Adiciona item com marcador estilizado"""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    
    run_bold = p.add_run(bold_prefix)
    run_bold.font.name = "Arial"
    run_bold.font.bold = True
    run_bold.font.size = Pt(10)
    run_bold.font.color.rgb = RGBColor(15, 23, 42)
    
    run_txt = p.add_run(text)
    run_txt.font.name = "Arial"
    run_txt.font.size = Pt(10)
    run_txt.font.color.rgb = RGBColor(51, 65, 85)

print("Iniciando montagem do E-Book DOCX...")
doc = Document()

# Configuração de Margens
for sec in doc.sections:
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)

# ==============================================================================
# CAPA DO E-BOOK
# ==============================================================================
p_espaco = doc.add_paragraph()
p_espaco.paragraph_format.space_before = Pt(40)

p_tag = doc.add_paragraph()
p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_tag = p_tag.add_run("GUIA DEFINITIVO DO USUÁRIO & MANUAL DE OPERAÇÕES")
r_tag.font.name = "Arial"
r_tag.font.size = Pt(11)
r_tag.font.bold = True
r_tag.font.color.rgb = RGBColor(37, 99, 235) # Blue 600

p_titulo = doc.add_paragraph()
p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_titulo.paragraph_format.space_before = Pt(10)
p_titulo.paragraph_format.space_after = Pt(10)
r_titulo = p_titulo.add_run("PAINEL DE GESTÃO PROFISSIONAL")
r_titulo.font.name = "Arial"
r_titulo.font.size = Pt(28)
r_titulo.font.bold = True
r_titulo.font.color.rgb = RGBColor(15, 23, 42) # Slate 900

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_after = Pt(25)
r_sub = p_sub.add_run("Domine todas as funcionalidades: Ordens de Serviço, Vendas, Controle de Estoque, Caixa Financeiro, Comissões de Vendedores e Mecânicos, Rastreamento Online e Muito Mais.")
r_sub.font.name = "Arial"
r_sub.font.size = Pt(12)
r_sub.font.color.rgb = RGBColor(100, 116, 139) # Slate 500

# Faixa decorativa
tbl_faixa = doc.add_table(rows=1, cols=1)
tbl_faixa.alignment = WD_TABLE_ALIGNMENT.CENTER
c_faixa = tbl_faixa.cell(0, 0)
c_faixa.width = Inches(6.5)
set_cell_background(c_faixa, "1E293B")
set_cell_margins(c_faixa, top=140, bottom=140, left=180, right=180)
p_faixa = c_faixa.paragraphs[0]
p_faixa.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_faixa = p_faixa.add_run("OFICINAS MECÂNICAS • AUTO CENTERS • LOJAS • COMÉRCIO EM GERAL • PRESTADORES DE SERVIÇOS")
r_faixa.font.name = "Arial"
r_faixa.font.size = Pt(9.5)
r_faixa.font.bold = True
r_faixa.font.color.rgb = RGBColor(255, 255, 255)

p_espaco2 = doc.add_paragraph()
p_espaco2.paragraph_format.space_before = Pt(120)

p_autor = doc.add_paragraph()
p_autor.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_autor.paragraph_format.space_after = Pt(2)
r_autor = p_autor.add_run("Desenvolvido para Simplificar e Alavancar seu Negócio")
r_autor.font.name = "Arial"
r_autor.font.size = Pt(10)
r_autor.font.bold = True
r_autor.font.color.rgb = RGBColor(51, 65, 85)

p_versao = doc.add_paragraph()
p_versao.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_versao = p_versao.add_run("Edição Oficial • Atualizada com Múltiplas Formas de Pagamento e Gestão de Comissões")
r_versao.font.name = "Arial"
r_versao.font.size = Pt(9)
r_versao.font.color.rgb = RGBColor(148, 163, 184)

doc.add_page_break()

# ==============================================================================
# SUMÁRIO
# ==============================================================================
p_sum = doc.add_paragraph()
format_heading(p_sum, "Sumário do E-Book", 1)

sumario_itens = [
    ("Capítulo 1", "Visão Geral e Primeiros Passos no Painel"),
    ("Capítulo 2", "Suporte Multi-Segmento: Modo OFICINA x Modo LOJA"),
    ("Capítulo 3", "Cadastro de Clientes e Gestão Veicular"),
    ("Capítulo 4", "Ordens de Serviço e Pedidos: O Coração Operacional"),
    ("Capítulo 5", "Equipe, Vendedores e Mecânicos: Controle Completo de Comissões"),
    ("Capítulo 6", "Múltiplas Formas de Pagamento: Flexibilidade Total no Fechamento"),
    ("Capítulo 7", "Experiência do Cliente: Rastreamento em Tempo Real & Assinatura Digital"),
    ("Capítulo 8", "Impressões Profissionais: Cupom Térmico 80mm & Ordem em PDF A4"),
    ("Capítulo 9", "Controle de Estoque e Catálogo de Produtos e Serviços"),
    ("Capítulo 10", "Gestão Financeira e Livro Caixa Integrado"),
    ("Capítulo 11", "Retornos Preventivos: Como Fidelizar Clientes Automaticamente"),
    ("Capítulo 12", "Módulo Fiscal (Opcional): Emissão de NFe/NFSe e Central Contábil"),
    ("Capítulo 13", "Boas Práticas e Dicas de Sucesso para seu Negócio"),
]

tbl_sum = doc.add_table(rows=0, cols=2)
tbl_sum.alignment = WD_TABLE_ALIGNMENT.CENTER
for cap, desc in sumario_itens:
    row = tbl_sum.add_row()
    c1, c2 = row.cells
    c1.width = Inches(1.8)
    c2.width = Inches(4.7)
    set_cell_margins(c1, top=50, bottom=50, left=60, right=60)
    set_cell_margins(c2, top=50, bottom=50, left=60, right=60)
    
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_before = Pt(2)
    p1.paragraph_format.space_after = Pt(2)
    r1 = p1.add_run(cap)
    r1.font.name = "Arial"
    r1.font.bold = True
    r1.font.size = Pt(10)
    r1.font.color.rgb = RGBColor(37, 99, 235)
    
    p2 = c2.paragraphs[0]
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run(desc)
    r2.font.name = "Arial"
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor(51, 65, 85)

doc.add_page_break()

# ==============================================================================
# CAPÍTULO 1
# ==============================================================================
p1 = doc.add_paragraph()
format_heading(p1, "Capítulo 1: Visão Geral e Primeiros Passos no Painel", 1)

add_body_p(doc, "O Painel de Gestão foi desenvolvido para ser a ferramenta central e definitiva do seu negócio. Ele elimina o uso de planilhas confusas, cadernos de papel e controles paralelos, reunindo em uma única plataforma na nuvem todo o fluxo de atendimento, execução de serviços, controle financeiro e fidelização de clientes.")

add_body_p(doc, "O sistema é 100% responsivo, funcionando com a mesma fluidez no computador, tablet ou no smartphone do dono e da sua equipe.")

format_heading(doc.add_paragraph(), "1.1 Como Acessar e Configurar sua Empresa", 2)
add_body_p(doc, "Assim que você realiza o login pela primeira vez, o primeiro passo recomendado é configurar a identidade visual e os dados cadastrais do seu estabelecimento:")

add_bullet_item(doc, "Menu Configurações (Engrenagem no topo): ", "Acesse para preencher o Nome Fantasia, Razão Social, CNPJ/CPF, Telefone de Contato e WhatsApp Comercial.")
add_bullet_item(doc, "Endereço Completo: ", "Informe a rua, número, bairro e cidade. Esses dados saem automaticamente impressos no cabeçalho dos comprovantes e na Ordem de Serviço.")
add_bullet_item(doc, "Logomarca Própria: ", "Faça o upload da sua logomarca (em formato PNG, JPG ou SVG). O sistema salvará e exibirá a sua marca no topo da tela, nos cupons e em todas as vias em PDF geradas para seus clientes.")
add_bullet_item(doc, "Mensagem de Rodapé: ", "Personalize a mensagem final das suas impressões (Exemplo: 'Garantia de 90 dias em todas as peças e serviços. Agradecemos a preferência!').")

add_callout(doc, "Se você ainda não tiver uma imagem de logotipo pronta, não se preocupe: o sistema exibe automaticamente um monograma corporativo moderno com as iniciais da sua empresa, mantendo uma aparência 100% profissional desde o primeiro dia.", "DICA DE APRESENTAÇÃO", "tip")

format_heading(doc.add_paragraph(), "1.2 Conhecendo o Dashboard Principal", 2)
add_body_p(doc, "Ao acessar o painel, a tela principal apresenta os principais indicadores de saúde da sua empresa em tempo real:")
add_bullet_item(doc, "Faturamento do Dia: ", "Soma de todas as entradas e pagamentos recebidos na data de hoje.")
add_bullet_item(doc, "Despesas do Dia: ", "Saídas registradas no caixa hoje para controle do saldo líquido diário.")
add_bullet_item(doc, "Faturamento Acumulado do Mês: ", "Total recebido no mês corrente menos despesas, permitindo acompanhar o crescimento da empresa.")
add_bullet_item(doc, "Tabela de Ordens em Aberto: ", "Lista imediata de todos os serviços que estão em andamento na loja ou oficina, mostrando cliente, valor e status.")

# ==============================================================================
# CAPÍTULO 2
# ==============================================================================
doc.add_page_break()
p2 = doc.add_paragraph()
format_heading(p2, "Capítulo 2: Suporte Multi-Segmento: Modo OFICINA x Modo LOJA", 1)

add_body_p(doc, "Uma das maiores inovações do Painel de Gestão é a capacidade de se adaptar dinamicamente ao segmento exato da sua empresa com apenas um clique nas Configurações.")

format_heading(doc.add_paragraph(), "2.1 Entendendo as Diferenças de Modo", 2)

tbl_seg = doc.add_table(rows=1, cols=3)
tbl_seg.alignment = WD_TABLE_ALIGNMENT.CENTER
h1, h2, h3 = tbl_seg.rows[0].cells
h1.width = Inches(1.8)
h2.width = Inches(2.3)
h3.width = Inches(2.4)
set_cell_background(h1, "1E293B")
set_cell_background(h2, "1E293B")
set_cell_background(h3, "1E293B")
for c, txt in [(h1, "RECURSO"), (h2, "MODO OFICINA MECÂNICA"), (h3, "MODO LOJA / COMÉRCIO")]:
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt)
    r.font.name = "Arial"
    r.font.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(255, 255, 255)

dados_seg = [
    ("Terminologia Principal", "Ordem de Serviço (OS)", "Venda / Pedido"),
    ("Identificação Veicular", "Placa, Marca, Modelo e Ano", "Oculto (Foco no cliente e produtos)"),
    ("Profissional Executor", "Mecânico / Técnico Responsável", "Vendedor / Atendente"),
    ("Status de Conclusão", "'Veículo Pronto para Retirada'", "'Pedido Pronto para Entrega'"),
    ("Garantia & Retorno", "Revisão preventiva (Km / Dias)", "Garantia comercial e pós-venda")
]

for rec, ofi, loj in dados_seg:
    row = tbl_seg.add_row()
    c1, c2, c3 = row.cells
    c1.width = Inches(1.8)
    c2.width = Inches(2.3)
    c3.width = Inches(2.4)
    set_cell_margins(c1, top=60, bottom=60, left=80, right=80)
    set_cell_margins(c2, top=60, bottom=60, left=80, right=80)
    set_cell_margins(c3, top=60, bottom=60, left=80, right=80)
    
    for c, val, bold in [(c1, rec, True), (c2, ofi, False), (c3, loj, False)]:
        p = c.paragraphs[0]
        r = p.add_run(val)
        r.font.name = "Arial"
        r.font.size = Pt(9)
        r.font.bold = bold
        r.font.color.rgb = RGBColor(51, 65, 85)

add_callout(doc, "Você pode alterar o tipo de negócio a qualquer momento em 'Configurações da Empresa'. O sistema ajusta automaticamente todos os formulários e termos sem perder nenhum dado já registrado.", "MUDANÇA DE MODO SEGURA", "success")

# ==============================================================================
# CAPÍTULO 3
# ==============================================================================
doc.add_page_break()
p3 = doc.add_paragraph()
format_heading(p3, "Capítulo 3: Cadastro de Clientes e Gestão Veicular", 1)

add_body_p(doc, "Ter uma base de clientes organizada é o primeiro pilar para vendas recorrentes. O Painel de Gestão permite cadastrar tanto pessoas físicas quanto pessoas jurídicas com facilidade.")

format_heading(doc.add_paragraph(), "3.1 Cadastro Completo do Cliente", 2)
add_bullet_item(doc, "Identificação: ", "Nome completo do cliente ou Razão Social da empresa.")
add_bullet_item(doc, "Telefone e WhatsApp: ", "Campo com formatação automática. Ao clicar no ícone do WhatsApp no painel, o sistema abre a conversa direto no WhatsApp do cliente com um clique.")
add_bullet_item(doc, "CPF / CNPJ: ", "Essencial para emissão de notas fiscais e identificação em pedidos faturados.")
add_bullet_item(doc, "Endereço com CEP: ", "Rua, Bairro, Cidade e UF para fins de entrega ou faturamento.")

format_heading(doc.add_paragraph(), "3.2 Gestão de Veículos (Modo Oficina)", 2)
add_body_p(doc, "No Modo Oficina, cada cliente pode ter um ou múltiplos veículos associados ao seu cadastro:")
add_bullet_item(doc, "Placa do Veículo: ", "Aceita o padrão Mercosul (ex: BRA2E19) ou padrão tradicional cinza (ex: ABC-1234).")
add_bullet_item(doc, "Marca e Modelo: ", "Identificação exata do carro ou moto (ex: VW Gol 1.6, Fiat Strada, Honda Civic).")
add_bullet_item(doc, "Histórico por Veículo: ", "Ao consultar o veículo, você tem acesso a todo o histórico de manutenções já realizadas, peças trocadas e datas de retorno.")

# ==============================================================================
# CAPÍTULO 4
# ==============================================================================
doc.add_page_break()
p4 = doc.add_paragraph()
format_heading(p4, "Capítulo 4: Ordens de Serviço e Pedidos: O Coração Operacional", 1)

add_body_p(doc, "A Ordem de Serviço (OS) ou Pedido de Venda é onde tudo acontece. Foi desenhada para ser rápida de abrir no balcão e completa ao finalizar.")

format_heading(doc.add_paragraph(), "4.1 Passo a Passo: Abrindo uma Nova OS / Pedido", 2)
add_bullet_item(doc, "1. Clique em 'Nova OS' ou 'Novo Pedido': ", "Localizado no topo do sistema com botão de destaque.")
add_bullet_item(doc, "2. Selecione o Cliente: ", "Comece a digitar o nome do cliente. Se ele já existir, o sistema completa os dados instantaneamente. Se for novo, cadastre na mesma hora sem sair da tela.")
add_bullet_item(doc, "3. Identifique o Veículo (Oficina): ", "Escolha a placa ou adicione um novo carro do cliente.")
add_bullet_item(doc, "4. Descreva a Solicitação / Problema: ", "Relate o problema informado pelo cliente (ex: 'Barulho na suspensão dianteira ao passar em lombadas' ou 'Solicitação de orçamento de tintas').")
add_bullet_item(doc, "5. Numeração Automática Isolada: ", "O sistema gera o número sequencial da sua empresa (ex: #0001, #0002), garantindo organização contábil impecável.")

format_heading(doc.add_paragraph(), "4.2 As 5 Etapas de Andamento em Tempo Real", 2)
add_body_p(doc, "Para dar clareza total ao cliente e à sua equipe, cada ordem passa por 5 etapas visuais:")
add_bullet_item(doc, "1. RECEBIDO (Azul): ", "O item/veículo acabou de dar entrada no estabelecimento.")
add_bullet_item(doc, "2. DIAGNÓSTICO (Âmbar): ", "Em fase de orçamento, testes técnicos ou conferência.")
add_bullet_item(doc, "3. EM EXECUÇÃO (Roxo): ", "Serviço aprovado e sendo realizado na bancada ou oficina.")
add_bullet_item(doc, "4. PRONTO (Verde): ", "Tudo finalizado, pronto para retirada ou entrega.")
add_bullet_item(doc, "5. ENTREGUE (Cinza): ", "Entregue nas mãos do cliente com pagamento liquidado.")

# ==============================================================================
# CAPÍTULO 5
# ==============================================================================
doc.add_page_break()
p5 = doc.add_paragraph()
format_heading(p5, "Capítulo 5: Equipe, Vendedores e Mecânicos: Controle Completo de Comissões", 1)

add_body_p(doc, "Muitas empresas pagam comissão aos vendedores pelo valor total vendido e aos mecânicos/técnicos pela mão de obra executada. O Painel de Gestão automatiza 100% dessa rotina complexa, eliminando erros e disputas.")

format_heading(doc.add_paragraph(), "5.1 Cadastro de Colaboradores", 2)
add_body_p(doc, "No menu 'Colaboradores', você cadastra cada membro da sua equipe:")
add_bullet_item(doc, "Função do Colaborador: ", "Defina se ele é 'Vendedor', 'Mecânico / Técnico' ou 'Ambos'.")
add_bullet_item(doc, "Chave Pix do Profissional: ", "Cadastre a chave Pix (CPF, celular, e-mail) para facilitar os acertos semanais ou mensais.")
add_bullet_item(doc, "Porcentagem Padrão de Comissão (%): ", "Informe o percentual que esse profissional costuma receber (ex: 5% para vendedor, 40% para mecânico).")
add_bullet_item(doc, "Base de Cálculo: ", "Escolha se a comissão incide sobre o Total Geral da Venda, apenas sobre Peças ou apenas sobre Mão de Obra.")

format_heading(doc.add_paragraph(), "5.2 Vinculação e Ajuste Flexível na Conclusão da OS", 2)
add_body_p(doc, "Ao finalizar uma Ordem de Serviço ou Pedido, você encontrará o bloco 'Equipe Responsável & Comissões':")
add_bullet_item(doc, "Seleção do Vendedor: ", "Escolha quem vendeu o serviço. O sistema puxa a porcentagem padrão dele, mas você é livre para alterar a porcentagem para mais ou para menos nessa OS específica!")
add_bullet_item(doc, "Seleção do Mecânico/Técnico: ", "Escolha quem executou a mão de obra. O sistema calcula a comissão em tempo real sobre os serviços realizados.")
add_bullet_item(doc, "Cálculo Imediato na Tela: ", "Conforme você altera o percentual, o valor em Reais (R$) a pagar ao colaborador atualiza na hora.")

format_heading(doc.add_paragraph(), "5.3 Relatório Inteligente: Comissões por Dia, Semana e Mês", 2)
add_body_p(doc, "No menu 'Comissões', o dono do estabelecimento tem o raio-x completo para realizar os pagamentos:")
add_bullet_item(doc, "Filtro Rápido 'Hoje': ", "Veja tudo o que foi gerado no dia de trabalho.")
add_bullet_item(doc, "Filtro Rápido 'Esta Semana': ", "Ideal para oficinas e lojas que fecham o pagamento aos sábados.")
add_bullet_item(doc, "Filtro Rápido 'Este Mês': ", "Perfeito para fechamento mensal junto com a folha de pagamento.")
add_bullet_item(doc, "Botão Copiar Chave Pix: ", "Ao lado do nome de cada colaborador, há um botão para copiar a chave Pix dele e abrir o aplicativo do seu banco para transferir o valor em segundos.")

add_callout(doc, "O dono do negócio tem total liberdade: se uma determinada OS não tiver comissão, basta deixar em branco ou colocar 0%. O sistema respeita integralmente a sua regra de negócio.", "FLEXIBILIDADE DE GESTÃO", "tip")

# ==============================================================================
# CAPÍTULO 6
# ==============================================================================
doc.add_page_break()
p6 = doc.add_paragraph()
format_heading(p6, "Capítulo 6: Múltiplas Formas de Pagamento: Flexibilidade Total no Fechamento", 1)

add_body_p(doc, "No dia a dia do comércio e das oficinas, é extremamente comum o cliente querer dividir o pagamento. Por exemplo: pagar uma parte em dinheiro que tem na carteira, outra parte no cartão de crédito e o restante no Pix. O Painel de Gestão oferece suporte nativo e visual a pagamentos mistos.")

format_heading(doc.add_paragraph(), "6.1 Alternador 'Única' vs 'Múltiplas (Dividir)'", 2)
add_body_p(doc, "Na hora de concluir o serviço ou registrar o pagamento, você tem duas opções:")
add_bullet_item(doc, "Modo 'Única': ", "Para pagamentos tradicionais 100% em uma única forma (Ex: 100% Pix, 100% Cartão ou 100% Dinheiro).")
add_bullet_item(doc, "Modo 'Múltiplas (Dividir)': ", "Permite adicionar quantas linhas de pagamento forem necessárias com formas e valores diferentes.")

format_heading(doc.add_paragraph(), "6.2 Como Funciona o Pagamento Misto na Prática", 2)
add_body_p(doc, "Imagine uma Ordem de Serviço no valor total de R$ 100,00:")
add_bullet_item(doc, "Linha 1: ", "Forma: Dinheiro | Valor: R$ 20,00")
add_bullet_item(doc, "Linha 2: ", "Forma: Cartão de Crédito | Valor: R$ 30,00")
add_bullet_item(doc, "Linha 3: ", "Forma: Pix | Valor: R$ 50,00")

add_body_p(doc, "Conforme você digita os valores, um indicador inteligente mostra:")
add_bullet_item(doc, "Total da OS: ", "R$ 100,00 (referência fixa).")
add_bullet_item(doc, "Status de Cobertura: ", "Se faltar algum centavo, o sistema mostra '⚠️ Falta: R$ XX,XX'. Se fechar certinho, exibe '✅ 100% Coberto (R$ 100,00)'. Se o cliente der dinheiro a mais, mostra 'ℹ️ Excedente / Troco: R$ XX,XX'.")

format_heading(doc.add_paragraph(), "6.3 Integração Perfeita com o Livro Caixa", 2)
add_body_p(doc, "Ao concluir uma OS com pagamento misto, o sistema não lança um valor genérico no seu financeiro. Ele gera movimentações individuais no Livro Caixa:")
add_bullet_item(doc, "Entrada 1: ", "Recebimento OS #0012 (Dinheiro) - R$ 20,00")
add_bullet_item(doc, "Entrada 2: ", "Recebimento OS #0012 (Cartão de Crédito) - R$ 30,00")
add_bullet_item(doc, "Entrada 3: ", "Recebimento OS #0012 (Pix) - R$ 50,00")
add_body_p(doc, "Dessa forma, a conciliação do seu caixa físico e das contas bancárias fica 100% exata ao final do dia!")

format_heading(doc.add_paragraph(), "6.4 Reedição de Pagamento em OS Concluída", 2)
add_body_p(doc, "O cliente pagou no dinheiro, mas depois pediu para trocar para cartão? Não tem problema: em qualquer OS já concluída, há um botão 'Alterar Forma de Pagamento'. Você pode reabrir o modal, ajustar as formas e o caixa é recalculado automaticamente.")

# ==============================================================================
# CAPÍTULO 7
# ==============================================================================
doc.add_page_break()
p7 = doc.add_paragraph()
format_heading(p7, "Capítulo 7: Experiência do Cliente: Rastreamento em Tempo Real & Assinatura Digital", 1)

add_body_p(doc, "Para destacar o seu negócio da concorrência e transmitir uma imagem de alta tecnologia, o Painel de Gestão conta com duas ferramentas exclusivas:")

format_heading(doc.add_paragraph(), "7.1 Rastreamento Online da Ordem de Serviço (Link Público)", 2)
add_body_p(doc, "Cada Ordem de Serviço gerada no sistema recebe um código de rastreio único e seguro. Você pode enviar o link direto para o WhatsApp do cliente com apenas um clique no botão verde 'Enviar WhatsApp'.")
add_body_p(doc, "Ao clicar no link pelo celular, o seu cliente visualiza:")
add_bullet_item(doc, "Logotipo e Contatos da sua Empresa: ", "Reforçando a credibilidade da sua marca.")
add_bullet_item(doc, "Barra de Progresso Visual: ", "O cliente vê se o pedido está em Diagnóstico, em Execução ou Pronto.")
add_bullet_item(doc, "Lista de Peças e Serviços: ", "Transparência total dos itens orçados e executados.")
add_bullet_item(doc, "Comprovante de Pagamento: ", "Discriminação clara de como o pedido foi pago.")
add_bullet_item(doc, "Termos de Garantia: ", "Orientando o cliente sobre os prazos de cobertura.")

format_heading(doc.add_paragraph(), "7.2 Assinatura Digital na Tela (Sem Papel)", 2)
add_body_p(doc, "Chega de imprimir folhas apenas para colher a rubrica do cliente. Na própria tela da Ordem de Serviço, há um painel de assinatura digital:")
add_bullet_item(doc, "Assinatura com o Dedo ou Caneta Touch: ", "O cliente assina diretamente na tela do seu tablet ou celular.")
add_bullet_item(doc, "Data e Hora Automáticas: ", "O sistema carimba a data e hora exatas em que a assinatura foi colhida.")
add_bullet_item(doc, "Validade Jurídica e Proteção: ", "A assinatura fica salva permanentemente no histórico da OS e sai impressa em todas as vias em PDF.")

# ==============================================================================
# CAPÍTULO 8
# ==============================================================================
doc.add_page_break()
p8 = doc.add_paragraph()
format_heading(p8, "Capítulo 8: Impressões Profissionais: Cupom Térmico 80mm & Ordem em PDF A4", 1)

add_body_p(doc, "O sistema foi projetado para atender todos os tipos de equipamentos e preferências de impressão do balcão:")

format_heading(doc.add_paragraph(), "8.1 Cupom Térmico Não-Fiscal (Impressoras 80mm / 58mm)", 2)
add_body_p(doc, "Ideal para agilidade no balcão e economia de papel. Ao clicar em 'Imprimir Cupom':")
add_bullet_item(doc, "Design Otimizado: ", "Largura perfeita para impressoras térmicas (Epson, Bematech, Elgin, POS).")
add_bullet_item(doc, "Detalhamento Limpo: ", "Contém dados da empresa, número da OS, cliente, itens resumidos, formas de pagamento discriminadas e rodapé.")
add_bullet_item(doc, "Auto-Impressão: ", "Abre o diálogo de impressão instantaneamente.")

format_heading(doc.add_paragraph(), "8.2 Ordem de Serviço Completa em A4 / PDF", 2)
add_body_p(doc, "Para entregas formais, faturamentos corporativos ou clientes que exigem relatório completo:")
add_bullet_item(doc, "Cabeçalho Oficial: ", "Logomarca em alta definição, dados fiscais da sua empresa e dados cadastrais do cliente.")
add_bullet_item(doc, "Tabela Detalhada de Itens: ", "Discriminação de Peças/Produtos com quantidade e valores, e Mão de Obra/Serviços com subtotais.")
add_bullet_item(doc, "Quadro de Pagamento e Garantia: ", "Composição exata dos pagamentos e prazos de retorno.")
add_bullet_item(doc, "Campo de Assinatura: ", "Exibe a assinatura digital colhida na tela ou campo pontilhado para assinatura física.")

# ==============================================================================
# CAPÍTULO 9
# ==============================================================================
doc.add_page_break()
p9 = doc.add_paragraph()
format_heading(p9, "Capítulo 9: Controle de Estoque e Catálogo de Produtos e Serviços", 1)

add_body_p(doc, "Saber exatamente o que você tem na prateleira e a margem de lucro de cada item é vital para não perder dinheiro.")

format_heading(doc.add_paragraph(), "9.1 Cadastro de Peças e Produtos", 2)
add_bullet_item(doc, "Código de Barras / Referência: ", "Facilita a busca com leitor de código de barras ou código do fabricante.")
add_bullet_item(doc, "Preço de Custo e Preço de Venda: ", "O sistema exibe a margem de lucro estimada.")
add_bullet_item(doc, "Estoque Atual e Estoque Mínimo: ", "Quando a quantidade de um item atinge o nível mínimo, o sistema destaca visualmente para alertar a necessidade de reposição.")
add_bullet_item(doc, "Baixa Automática no Estoque: ", "Ao vincular um produto na Ordem de Serviço e concluí-la, a quantidade é deduzida automaticamente do estoque.")

format_heading(doc.add_paragraph(), "9.2 Catálogo de Serviços Padronizados", 2)
add_body_p(doc, "Cadastre seus serviços mais comuns (ex: 'Alinhamento e Balanceamento', 'Troca de Óleo', 'Instalação Elétrica', 'Formatação'). Na hora de abrir a OS, basta começar a digitar o nome do serviço para puxar o valor padrão instantaneamente!")

# ==============================================================================
# CAPÍTULO 10
# ==============================================================================
doc.add_page_break()
p10 = doc.add_paragraph()
format_heading(p10, "Capítulo 10: Gestão Financeira e Livro Caixa Integrado", 1)

add_body_p(doc, "O módulo de Financeiro mantém o controle rigoroso de tudo o que entra e sai da empresa:")

format_heading(doc.add_paragraph(), "10.1 Lançamentos Automáticos x Lançamentos Manuais", 2)
add_bullet_item(doc, "Lançamentos Automáticos de Receita: ", "Toda vez que uma OS ou Pedido é concluído, o sistema lança automaticamente a receita no caixa correspondente.")
add_bullet_item(doc, "Lançamentos Manuais de Despesas: ", "Registre pagamentos de água, luz, aluguel, compra de ferramentas, café e suprimentos.")
add_bullet_item(doc, "Lançamentos de Outras Receitas: ", "Recebeu algum valor avulso fora de OS? Lance rapidamente em 'Nova Receita'.")

format_heading(doc.add_paragraph(), "10.2 Fechamento e Saldo do Caixa", 2)
add_body_p(doc, "Ao final do dia, o sistema calcula:")
add_bullet_item(doc, "Total de Receitas do Dia: ", "Soma de todo o dinheiro, pix e cartões recebidos.")
add_bullet_item(doc, "Total de Despesas do Dia: ", "Soma das contas pagas hoje.")
add_bullet_item(doc, "Saldo Líquido: ", "A diferença exata entre receitas e despesas.")

# ==============================================================================
# CAPÍTULO 11
# ==============================================================================
doc.add_page_break()
p11 = doc.add_paragraph()
format_heading(p11, "Capítulo 11: Retornos Preventivos: Como Fidelizar Clientes Automaticamente", 1)

add_body_p(doc, "O custo para conquistar um cliente novo é até 5 vezes maior do que manter um cliente existente. O módulo de Retornos Preventivos foi criado exatamente para gerar faturamento recorrente para a sua empresa.")

format_heading(doc.add_paragraph(), "11.1 Como Configurar o Retorno ao Concluir o Pedido", 2)
add_body_p(doc, "Ao finalizar qualquer serviço, o sistema pergunta qual é o prazo previsto para o próximo retorno:")
add_bullet_item(doc, "Em 30 dias (1 mês): ", "Ideal para revisões rápidas ou conferência de peças novas.")
add_bullet_item(doc, "Em 60 dias (2 meses): ", "Acompanhamento de serviços médios.")
add_bullet_item(doc, "Em 90 dias (3 meses - Padrão): ", "Padrão de garantia legal e revisão preventiva.")
add_bullet_item(doc, "Em 180 dias ou 1 ano: ", "Trocas de óleo periódicas, alinhamento ou manutenções anuais.")
add_bullet_item(doc, "Data Personalizada: ", "Você escolhe o dia exato no calendário.")

format_heading(doc.add_paragraph(), "11.2 A Central de Retornos e Contato Ativo", 2)
add_body_p(doc, "No menu 'Retornos', o sistema lista todos os clientes cuja data de revisão preventiva está vencendo ou já venceu:")
add_bullet_item(doc, "Botão WhatsApp Direto: ", "Com um único clique, você abre o WhatsApp do cliente com uma mensagem simpática convidando-o para a revisão.")
add_bullet_item(doc, "Resultado Comprovado: ", "Empresas que utilizam o módulo de retornos aumentam seu faturamento em até 30% com clientes que voltam espontaneamente.")

# ==============================================================================
# CAPÍTULO 12
# ==============================================================================
doc.add_page_break()
p12 = doc.add_paragraph()
format_heading(p12, "Capítulo 12: Módulo Fiscal (Opcional): Emissão de NFe/NFSe e Central Contábil", 1)

add_body_p(doc, "Para as empresas que necessitam de emissão de documento fiscal eletrônico, o Painel de Gestão oferece um módulo integrado à Focus NFe:")

format_heading(doc.add_paragraph(), "12.1 Tipos de Documentos Suportados", 2)
add_bullet_item(doc, "NFS-e (Nota Fiscal de Serviços Eletrônica): ", "Para mão de obra de oficinas, consultorias e prestadores de serviços com a prefeitura.")
add_bullet_item(doc, "NF-e (Nota Fiscal Eletrônica - Modelo 55): ", "Para venda de produtos e peças com envio interestadual ou empresas.")
add_bullet_item(doc, "NFC-e (Nota Fiscal de Consumidor Eletrônica): ", "Cupom fiscal direto para o consumidor no balcão.")

format_heading(doc.add_paragraph(), "12.2 Fechamento Fiscal Mensal em 1 Clique (Central Contábil)", 2)
add_body_p(doc, "Chega de perder horas no final do mês procurando arquivos XML para enviar ao seu contador:")
add_bullet_item(doc, "Exportação em ZIP: ", "Basta selecionar o mês desejado e clicar em 'Baixar Fechamento Fiscal'. O sistema compacta todos os arquivos XML autorizados e gera uma planilha detalhada (.CSV) com totais e tributos.")
add_bullet_item(doc, "Envio ao Contador: ", "Você só precisa anexar esse arquivo ZIP no e-mail do seu escritório de contabilidade!")

# ==============================================================================
# CAPÍTULO 13
# ==============================================================================
doc.add_page_break()
p13 = doc.add_paragraph()
format_heading(p13, "Capítulo 13: Boas Práticas e Dicas de Sucesso para seu Negócio", 1)

add_body_p(doc, "Para extrair 100% do potencial do Painel de Gestão e levar sua empresa ao próximo nível, siga este checklist de boas práticas diárias:")

add_bullet_item(doc, "1. Abra a OS Imediatamente: ", "Nunca comece um serviço sem abrir a OS no painel. Isso protege a sua empresa com registro claro do que foi solicitado.")
add_bullet_item(doc, "2. Colete a Assinatura na Tela: ", "Acostume seus clientes a assinar no tablet ou celular no momento da entrada ou da entrega. Isso traz respaldo jurídico total.")
add_bullet_item(doc, "3. Mantenha os Telefones com WhatsApp Atualizados: ", "O WhatsApp é a ponte de ouro para enviar o link de rastreio e os lembretes de retorno preventivo.")
add_bullet_item(doc, "4. Feche o Caixa Diariamente: ", "Ao final do expediente, confira se o total em dinheiro físico na gaveta confere com o relatório de receitas em dinheiro do painel.")
add_bullet_item(doc, "5. Acerte as Comissões com Transparência: ", "Use o módulo de Comissões para mostrar ao seu colaborador exatamente quais OSs compuseram o bônus dele. Isso motiva a equipe e gera confiança.")

add_callout(doc, "O Painel de Gestão é atualizado continuamente com melhorias de performance, novas ferramentas e segurança de ponta na nuvem para apoiar o crescimento da sua empresa!", "SUCESSO NO SEU NEGÓCIO", "success")

# Salvar o documento final
doc.save(DOCX_OUTPUT)
print(f"[SUCESSO TOTAL] E-Book gerado com sucesso em: {DOCX_OUTPUT}")
